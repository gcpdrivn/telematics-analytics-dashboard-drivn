"""Excel (.xlsx) generation for the vehicle uptime calendar. Cell coloring and
comments need per-cell control that pandas.to_excel doesn't give, so this
builds the workbook directly with openpyxl."""

from __future__ import annotations

import datetime as dt
import io

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from backend.uptime import COMBINED_COLORS, COMBINED_INDEX, LEGEND_COMBINED

FIXED_HEADERS = ["Vehicle Number", "Vehicle Type", "Vehicle Model", "Customer Name", "Uptime %"]


def _fill(hex_color: str) -> PatternFill:
    c = hex_color.lstrip("#").upper()
    return PatternFill(start_color=f"FF{c}", end_color=f"FF{c}", fill_type="solid")


def _contrast_text(hex_color: str) -> str:
    c = hex_color.lstrip("#")
    r, g, b = int(c[0:2], 16), int(c[2:4], 16), int(c[4:6], 16)
    luminance = (0.299 * r + 0.587 * g + 0.114 * b) / 255
    return "FF1A1A1A" if luminance > 0.6 else "FFFFFFFF"


def _pct_color(pct: float | None) -> str:
    if pct is None:
        return "FF808080"
    if pct >= 90:
        return "FF2F9E44"
    if pct >= 70:
        return "FFE8590C"
    return "FFE03131"


def _add_legend_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("Legend")
    ws.append(["Index", "Color", "Status", "Meaning"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r, item in enumerate(LEGEND_COMBINED, start=2):
        idx_cell = ws.cell(row=r, column=1, value=item["index"])
        idx_cell.alignment = Alignment(horizontal="center")
        idx_cell.font = Font(bold=True)
        swatch = ws.cell(row=r, column=2, value="")
        if item["color"]:
            swatch.fill = _fill(item["color"])
        ws.cell(row=r, column=3, value=item["status"])
        ws.cell(row=r, column=4, value=item["label"])
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 40


def build_uptime_workbook(vehicles: list[dict], dates_desc: list[str]) -> bytes:
    """dates_desc: ISO date strings, latest first (matches the column order
    the user asked for). vehicles: build_vehicle_uptime() output."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Uptime"

    ws.append(FIXED_HEADERS + [dt.date.fromisoformat(d).strftime("%d-%b-%y") for d in dates_desc])
    header_font = Font(bold=True)
    for cell in ws[1]:
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    ws["A1"].comment = Comment("Each day cell shows the legend index for its status -- see the Legend sheet.", "Drivn Uptime Report")

    for r, v in enumerate(vehicles, start=2):
        ws.cell(row=r, column=1, value=v["vehicle_number"])
        ws.cell(row=r, column=2, value=v["vehicle_type"])
        ws.cell(row=r, column=3, value=v["vehicle_model"])
        ws.cell(row=r, column=4, value=v["customer_name"])
        pct_cell = ws.cell(row=r, column=5, value=v["uptime_pct"])
        pct_cell.number_format = "0.0"
        pct_cell.font = Font(bold=True, color=_pct_color(v["uptime_pct"]))
        pct_cell.alignment = Alignment(horizontal="center")

        by_date = {day["date"]: day for day in v["daily"]}
        for c, d in enumerate(dates_desc, start=6):
            info = by_date.get(d)
            if info is None:
                continue
            status = info["combined_status"]
            color = COMBINED_COLORS[status]
            index = COMBINED_INDEX[status]
            cell = ws.cell(row=r, column=c, value=index)
            if color:
                cell.fill = _fill(color)
                cell.font = Font(bold=True, color=_contrast_text(color))
            else:
                cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")
            if info.get("note"):
                cell.comment = Comment(info["note"], "Drivn Uptime Report")

    ws.freeze_panes = "F2"
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 10
    for c in range(6, 6 + len(dates_desc)):
        ws.column_dimensions[get_column_letter(c)].width = 7

    _add_legend_sheet(wb)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
