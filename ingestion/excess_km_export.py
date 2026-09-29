"""Monthly excess-km workbook: `uv run export-excess-km --month 2026-08`.

Refreshes odometer_daily_resolved (so the month's readings are current),
then writes one Excel workbook for a completed month from the
excess_km_monthly view and its per-day audit rows:

- Summary        per customer, all SUMIFS/COUNTIFS over Vehicle Detail
- Vehicle Detail per vehicle: km, coverage, excess km and cost (formulas)
- Daily Readings every in-service day and how it was counted
- Terms          allowance / rate per customer and the review thresholds;
                 editing them here recalculates the workbook
- Method         how billable km are derived

Km measurements are values from BigQuery; everything derived from the
terms (excess km, cost, review flags) is an Excel formula.
"""

from __future__ import annotations

import datetime as dt
import logging
from pathlib import Path

import pandas as pd
from google.cloud import bigquery
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ingestion import bq_client, excess_km, odometer_resolver
from ingestion.config import REPO_ROOT, Settings

logger = logging.getLogger(__name__)

EXPORT_DIR = REPO_ROOT / "data" / "exports"

FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
TOTAL_FILL = PatternFill("solid", fgColor="D9E1F2")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
REVIEW_FILL = PatternFill("solid", fgColor="FCE4D6")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
# Indian digit grouping (1,23,456).
KM = '[>=10000000]##\\,##\\,##\\,##0;[>=100000]##\\,##\\,##0;##,##0'
INR = '[>=10000000]"₹"##\\,##\\,##\\,##0;[>=100000]"₹"##\\,##\\,##0;"₹"##,##0'
PCT = "0.0%"
RATIO = "0.000"

STATUS_LABELS = {
    "reliable": "Trusted reading",
    "bridged": "Silent day - trusted gap",
    "unverified": "Untrusted reading (not billed)",
    "rejected_gap": "Silent day - gap over limit (not billed)",
    "unbridged": "Silent day - no reading (not billed)",
}


class ExportError(RuntimeError):
    pass


def _query(client: bigquery.Client, sql: str, month: dt.date) -> pd.DataFrame:
    job = client.query(
        sql,
        job_config=bigquery.QueryJobConfig(
            query_parameters=[bigquery.ScalarQueryParameter("month", "DATE", month)]
        ),
    )
    return pd.DataFrame([dict(r) for r in job.result()])


def _font(cell, **kw) -> None:
    cell.font = Font(name=FONT, **kw)


def _header(ws, row: int, headers: list[str]) -> None:
    for col, text in enumerate(headers, 1):
        c = ws.cell(row=row, column=col, value=text)
        _font(c, bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    ws.row_dimensions[row].height = 45


def _widths(ws, widths: list[float]) -> None:
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _body(ws, first_row: int, last_row: int, n_cols: int, formats: dict[int, str]) -> None:
    for r in range(first_row, last_row + 1):
        for col in range(1, n_cols + 1):
            c = ws.cell(row=r, column=col)
            if c.font is None or c.font.name != FONT:
                _font(c)
            c.border = BORDER
            if col in formats:
                c.number_format = formats[col]


def _terms_sheet(wb: Workbook, customers: list[str], terms: dict) -> dict[str, str]:
    ws = wb.create_sheet("Terms")
    ws["A1"] = "Contract terms and review thresholds"
    _font(ws["A1"], bold=True, size=13)
    ws["A2"] = ("Yellow cells are inputs: change them and the workbook recalculates. "
                "Source: 'Excess KM & Battery Replacement.xlsx' (as loaded into dim_vehicle).")
    _font(ws["A2"], italic=True, color="595959")
    _header(ws, 4, ["Customer", "Monthly available km (per vehicle)", "Excess km rate (₹/km)"])
    for i, cust in enumerate(customers):
        r = 5 + i
        ws.cell(row=r, column=1, value=cust)
        for col, key in ((2, "monthly_available_km"), (3, "excess_km_rate")):
            c = ws.cell(row=r, column=col, value=terms[cust][key])
            c.fill = INPUT_FILL
            _font(c, color="0000FF")
    last = 4 + len(customers)
    _body(ws, 5, last, 3, {2: "#,##0.00", 3: '"₹"0.00'})

    t = last + 2
    _header(ws, t, ["Review threshold", "Value", "Meaning"])
    rows = [
        ("Minimum reliable coverage", excess_km.MIN_RELIABLE_COVERAGE, PCT,
         "Share of in-service days backed by a trusted reading or trusted silent gap; below this the vehicle is flagged."),
        ("Max odometer vs reported distance divergence", excess_km.MAX_ODOMETER_DISTANCE_DIVERGENCE, PCT,
         "Flag when odometer km differ from reported Distance on the same days by more than this."),
        ("Silent-gap limit (km/day)", excess_km.MAX_GAP_KM_PER_DAY, "#,##0",
         "Applied upstream: silent-stretch km averaging this or more per day were not billed. Informational here."),
    ]
    refs = {}
    for i, (label, val, fmt, meaning) in enumerate(rows):
        r = t + 1 + i
        ws.cell(row=r, column=1, value=label)
        c = ws.cell(row=r, column=2, value=val)
        c.number_format = fmt
        ws.cell(row=r, column=3, value=meaning)
        if i < 2:
            c.fill = INPUT_FILL
            _font(c, color="0000FF")
    _body(ws, t + 1, t + len(rows), 3, {})
    for i in range(3):
        ws.cell(row=t + 1 + i, column=3).alignment = Alignment(wrap_text=True, vertical="top")
    refs["coverage"] = f"Terms!$B${t + 1}"
    refs["divergence"] = f"Terms!$B${t + 2}"
    refs["cust_range"] = f"Terms!$A$5:$A${last}"
    refs["km_range"] = f"Terms!$B$5:$B${last}"
    refs["rate_range"] = f"Terms!$C$5:$C${last}"
    _widths(ws, [44, 22, 90])
    ws.freeze_panes = "A5"
    return refs


VEHICLE_HEADERS = [
    "Customer", "Vehicle", "Days in service", "Partial month",
    "Trusted reading days", "Trusted silent days", "Untrusted reading days",
    "Silent days over gap limit", "Silent days, no reading",
    "Reliable coverage", "Odometer km (trusted days)", "Silent-gap km (trusted)",
    "Billable km", "Reported distance km (same trusted days)", "Odometer ÷ reported",
    "Untrusted km (not billed)", "Monthly available km", "Excess km",
    "Excess km rate (₹/km)", "Excess cost (₹)", "Needs review", "Review reason",
]


def _vehicle_sheet(wb: Workbook, df: pd.DataFrame, refs: dict) -> int:
    ws = wb.create_sheet("Vehicle Detail")
    _header(ws, 1, VEHICLE_HEADERS)
    for i, v in enumerate(df.itertuples(), start=2):
        r = i
        values = [
            v.customer_name, v.base_license_plate, v.days_in_service,
            "Yes" if v.is_partial_month else "No",
            v.reliable_days, v.bridged_days, v.unverified_days, v.rejected_gap_days, v.unbridged_days,
        ]
        for col, val in enumerate(values, 1):
            ws.cell(row=r, column=col, value=val)
        ws.cell(row=r, column=10, value=f"=IF(C{r}=0,0,(E{r}+F{r})/C{r})")
        ws.cell(row=r, column=11, value=float(v.odometer_km))
        ws.cell(row=r, column=12, value=float(v.gap_km))
        ws.cell(row=r, column=13, value=f"=K{r}+L{r}")
        ws.cell(row=r, column=14, value=float(v.reported_km))
        ws.cell(row=r, column=15, value=f'=IF(N{r}=0,"",K{r}/N{r})')
        ws.cell(row=r, column=16, value=float(v.unverified_km))
        ws.cell(row=r, column=17, value=f"=INDEX({refs['km_range']},MATCH(A{r},{refs['cust_range']},0))")
        ws.cell(row=r, column=18, value=f"=MAX(0,M{r}-Q{r})")
        ws.cell(row=r, column=19, value=f"=INDEX({refs['rate_range']},MATCH(A{r},{refs['cust_range']},0))")
        ws.cell(row=r, column=20, value=f"=R{r}*S{r}")
        low = f"J{r}<{refs['coverage']}"
        mismatch = f'AND(O{r}<>"",ABS(O{r}-1)>{refs["divergence"]})'
        ws.cell(row=r, column=21, value=f'=IF(OR({low},{mismatch}),"Yes","No")')
        ws.cell(
            row=r, column=22,
            value=(f'=IF({low},"Low reliable coverage","")'
                   f'&IF(AND({low},{mismatch}),"; ","")'
                   f'&IF({mismatch},"Odometer vs reported distance mismatch","")'),
        )
    last = 1 + len(df)
    _body(ws, 2, last, len(VEHICLE_HEADERS), {
        10: PCT, 11: KM, 12: KM, 13: KM, 14: KM, 15: RATIO, 16: KM, 17: KM, 18: KM,
        19: '"₹"0.00', 20: INR,
    })
    for r in range(2, last + 1):
        _font(ws.cell(row=r, column=13), bold=True)
        _font(ws.cell(row=r, column=20), bold=True)
    ws.conditional_formatting.add(
        f"A2:{get_column_letter(len(VEHICLE_HEADERS))}{last}",
        FormulaRule(formula=[f'$U2="Yes"'], fill=REVIEW_FILL),
    )
    _widths(ws, [14, 14, 10, 9, 10, 10, 11, 11, 11, 11, 13, 12, 12, 15, 11, 12, 13, 11, 11, 13, 9, 40])
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(VEHICLE_HEADERS))}{last}"
    return last


def _summary_sheet(wb: Workbook, month: dt.date, customers: list[str], last_vehicle_row: int) -> None:
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = f"Excess km & cost - {month:%B %Y}"
    _font(ws["A1"], bold=True, size=14)
    ws["A2"] = (f"Per-vehicle allowance, calendar month {month:%d %b %Y} - "
                f"{(month + pd.offsets.MonthEnd(0)):%d %b %Y}. Generated {dt.date.today():%d %b %Y} "
                "from BigQuery (telematics.excess_km_monthly).")
    _font(ws["A2"], italic=True, color="595959")
    headers = ["Customer", "Vehicles", "Billable km", "Monthly available km (per vehicle)",
               "Vehicles over allowance", "Excess km", "Excess km rate (₹/km)", "Excess cost (₹)",
               "Vehicles needing review", "Excess cost from vehicles needing review (₹)"]
    _header(ws, 4, headers)
    vd = "'Vehicle Detail'!"
    rng = lambda col: f"{vd}${col}$2:${col}${last_vehicle_row}"
    for i, cust in enumerate(customers):
        r = 5 + i
        ws.cell(row=r, column=1, value=cust)
        crit = f"{rng('A')},$A{r}"
        ws.cell(row=r, column=2, value=f"=COUNTIFS({crit})")
        ws.cell(row=r, column=3, value=f"=SUMIFS({rng('M')},{crit})")
        ws.cell(row=r, column=4, value=f"=INDEX(Terms!$B:$B,MATCH($A{r},Terms!$A:$A,0))")
        ws.cell(row=r, column=5, value=f'=COUNTIFS({crit},{rng("R")},">0")')
        ws.cell(row=r, column=6, value=f"=SUMIFS({rng('R')},{crit})")
        ws.cell(row=r, column=7, value=f"=INDEX(Terms!$C:$C,MATCH($A{r},Terms!$A:$A,0))")
        ws.cell(row=r, column=8, value=f"=SUMIFS({rng('T')},{crit})")
        ws.cell(row=r, column=9, value=f'=COUNTIFS({crit},{rng("U")},"Yes")')
        ws.cell(row=r, column=10, value=f'=SUMIFS({rng("T")},{crit},{rng("U")},"Yes")')
    last = 4 + len(customers)
    tot = last + 1
    ws.cell(row=tot, column=1, value="Total")
    for col in (2, 3, 5, 6, 8, 9, 10):
        L = get_column_letter(col)
        ws.cell(row=tot, column=col, value=f"=SUM({L}5:{L}{last})")
    fmts = {3: KM, 4: KM, 6: KM, 7: '"₹"0.00', 8: INR, 10: INR}
    _body(ws, 5, tot, len(headers), fmts)
    for col in range(1, len(headers) + 1):
        c = ws.cell(row=tot, column=col)
        c.fill = TOTAL_FILL
        _font(c, bold=True)
    for r in range(5, last + 1):
        _font(ws.cell(row=r, column=8), bold=True)

    n = tot + 2
    notes = [
        "Notes",
        "- Excess is computed per vehicle: Excess km = MAX(0, Billable km - Monthly available km); Cost = Excess km x rate. "
        "One vehicle's unused km never offsets another's excess.",
        "- Billable km use trusted odometer readings only (see Method). Vehicles flagged 'Needs review' are still costed; "
        "check them on 'Vehicle Detail' (highlighted) before invoicing - a faulty odometer can also under-count, showing ₹0.",
        "- Vehicles that joined mid-month get the full monthly allowance (see 'Partial month' on Vehicle Detail).",
        "- BillionE and AVG LOGISTICS are not billed for excess km and are excluded.",
    ]
    for i, text in enumerate(notes):
        c = ws.cell(row=n + i, column=1, value=text)
        _font(c, bold=(i == 0))
    _widths(ws, [16, 10, 14, 16, 12, 13, 12, 15, 12, 18])
    ws.freeze_panes = "A5"


def _daily_sheet(wb: Workbook, daily: pd.DataFrame) -> None:
    ws = wb.create_sheet("Daily Readings")
    headers = ["Customer", "Vehicle", "Date", "How the day was counted", "Resolver fill method",
               "Opening odometer", "Closing odometer", "Odometer km (resolver)", "Reported distance km",
               "Billed odometer km", "Billed silent-gap km", "Billable km"]
    _header(ws, 1, headers)
    for i, d in enumerate(daily.itertuples(), start=2):
        vals = [
            d.customer_name, d.base_license_plate, d.day, STATUS_LABELS.get(d.status, d.status),
            d.fill_method, d.opening_odometer_clean, d.closing_odometer_clean,
            d.distance_by_odometer, d.distance_reported, d.odometer_km, d.gap_km,
        ]
        for col, val in enumerate(vals, 1):
            ws.cell(row=i, column=col, value=None if pd.isna(val) else val)
        ws.cell(row=i, column=12, value=f"=J{i}+K{i}")
    last = 1 + len(daily)
    _body(ws, 2, last, len(headers), {3: "dd-mmm-yyyy", 6: "#,##0.0", 7: "#,##0.0", 8: "#,##0.0",
                                      9: "#,##0.0", 10: "#,##0.0", 11: "#,##0.0", 12: "#,##0.0"})
    _widths(ws, [14, 14, 13, 36, 18, 14, 14, 13, 13, 12, 12, 12])
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{last}"


def _method_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("Method")
    lines = [
        ("How billable km are computed", True),
        ("1. Source: telematics.odometer_daily_resolved - the Fleetx daily odometer readings, with device faults "
         "(overflow values, resets, impossible jumps) detected and classified per day.", False),
        ("2. Trusted days: a day's odometer km count when the reading is a clean raw reading, or is filled between "
         "two clean readings (INTERPOLATED / MANUAL_OVERRIDE).", False),
        ("3. Silent days: when the device didn't report, the km between the last reading before and the first reading "
         f"after are counted (spread evenly over those days) only if they average under "
         f"{excess_km.MAX_GAP_KM_PER_DAY:,.0f} km/day.", False),
        ("4. Not billed: days whose reading is untrusted (fell back to the reported Distance, or unresolved), silent "
         "stretches over the km/day limit, and silent days with no later reading. These lower the vehicle's reliable coverage.", False),
        ("5. Excess km = MAX(0, Billable km - Monthly available km), per vehicle; Excess cost = Excess km x rate.", False),
        ("6. Needs review (flag only, still costed): reliable coverage below the Terms threshold, or odometer km on "
         "trusted days differing from the reported Distance on those same days by more than the Terms threshold.", False),
        ("", False),
        ("Day categories on 'Daily Readings'", True),
    ] + [(f"- {label}", False) for label in STATUS_LABELS.values()]
    for i, (text, bold) in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=text)
        _font(c, bold=bold, size=13 if bold and i == 1 else 10)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 120


def export(settings: Settings, month: dt.date, *, refresh: bool = True, out_dir: Path = EXPORT_DIR) -> Path:
    first_of_this_month = dt.date.today().replace(day=1)
    if month >= first_of_this_month:
        raise ExportError(f"{month:%B %Y} isn't complete yet -- only completed months can be exported.")
    client = bq_client.get_client(settings)
    if refresh:
        logger.info("Refreshing odometer_daily_resolved ...")
        odometer_resolver.run(settings)
    excess_km.ensure_views(client, settings)

    vehicles = _query(
        client,
        f"SELECT * FROM `{settings.excess_km_monthly_view_ref}` WHERE month = @month "
        "ORDER BY customer_name, base_license_plate",
        month,
    )
    if vehicles.empty:
        raise ExportError(f"No billed vehicles have readings in {month:%B %Y}.")
    daily = _query(client, excess_km.daily_rows_sql(settings) + " ORDER BY customer_name, base_license_plate, day", month)

    customers = sorted(vehicles["customer_name"].unique())
    terms = (
        vehicles.groupby("customer_name")[["monthly_available_km", "excess_km_rate"]].first().to_dict("index")
    )

    wb = Workbook()
    wb.active.title = "Summary"
    refs = _terms_sheet(wb, customers, terms)
    last_vehicle_row = _vehicle_sheet(wb, vehicles, refs)
    _summary_sheet(wb, month, customers, last_vehicle_row)
    _daily_sheet(wb, daily)
    _method_sheet(wb)
    wb.move_sheet("Terms", offset=len(wb.sheetnames))  # Summary, Vehicle Detail, Daily, Method, Terms
    wb.move_sheet("Method", offset=len(wb.sheetnames))

    for ws in wb.worksheets:
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True

    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"excess_km_{month:%Y-%m}.xlsx"
    wb.save(path)
    return path
