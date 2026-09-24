"""Cleaning logic for the manually-maintained dim_vehicle_master.xlsx --
per-vehicle OEM, type, model and telematics-device installation date, hand
filled in from the template seed_dimensions.py generates, plus an optional
Starting Odometer override.

Two independent things happen with this file:
- seed_dimensions.py reads it (via read_and_clean_vehicle_master) as the
  preferred source for oem/vehicle_type/vehicle_model/device_installation_date
  (and starting_odometer, when filled in) on every reseed, falling back to telemetry-derived values only for plates
  the file doesn't (yet) cover -- so re-running seed-dimensions never wipes
  out manually entered data.
- write_cleaned_copy() rewrites the sheet itself in place with the same
  cleaning applied (OEM's vehicle-type suffix stripped, install dates as real
  Excel dates in one consistent format), so the spreadsheet stays a clean
  reference too, not just whatever BigQuery ends up with.
"""

from __future__ import annotations

import datetime as dt
import re
from pathlib import Path

import pandas as pd

SHEET_NAME = "dim_vehicle master"

_COLUMN_RENAMES = {
    "Vehicle Number": "base_license_plate",
    "Customer Name": "customer_name",
    "OEM": "oem",
    "Vehicle Type": "vehicle_type",
    "Vehicle Model": "vehicle_model",
    "Device Installation Date": "device_installation_date",
}

# Optional columns: a sheet without them is still valid (read as blank).
# "Starting Odometer" is a manual override for dim_vehicle.starting_odometer,
# for a vehicle whose derived first reading is wrong -- e.g. DL1PD8677's new
# device reported 0.05 km for a day and a half before its odometer was set
# to the bus's real ~72,000 km.
_OPTIONAL_COLUMN_RENAMES = {
    "Starting Odometer": "starting_odometer",
}

FINAL_COLUMN_ORDER = [
    "base_license_plate",
    "customer_name",
    "oem",
    "vehicle_type",
    "vehicle_model",
    "device_installation_date",
    "starting_odometer",
]

# Strips a trailing "(...)" annotation some OEM values carry, e.g.
# "TATA (Truck)" -> "TATA" -- the vehicle type already has its own column.
_OEM_TYPE_SUFFIX_RE = re.compile(r"\s*\([^)]*\)\s*$")

INSTALL_DATE_NUMBER_FORMAT = "DD-MMM-YYYY"


def _clean_str(val) -> str | None:
    if val is None or (isinstance(val, float) and pd.isna(val)):
        return None
    s = str(val).strip()
    return s or None


def clean_oem(val) -> str | None:
    s = _clean_str(val)
    if s is None:
        return None
    return _OEM_TYPE_SUFFIX_RE.sub("", s).strip() or None


def _parse_install_date(val) -> dt.date | None:
    if val is None:
        return None
    try:
        if pd.isna(val):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(val, pd.Timestamp):
        return val.date()
    if isinstance(val, dt.datetime):
        return val.date()
    if isinstance(val, dt.date):
        return val
    s = str(val).strip()
    if not s:
        return None
    # Source cells are a mix of real Excel dates (already unambiguous) and
    # plain DD/MM/YYYY text (every value observed has a day component >12,
    # confirming day-first entry) -- dayfirst=True matches both.
    parsed = pd.to_datetime(s, dayfirst=True, errors="coerce")
    if pd.isna(parsed):
        raise ValueError(f"Could not parse Device Installation Date value: {val!r}")
    return parsed.date()


def _parse_odometer(val) -> float | None:
    s = _clean_str(val)
    if s is None:
        return None
    try:
        km = float(s.replace(",", ""))
    except ValueError as exc:
        raise ValueError(f"Could not parse Starting Odometer value: {val!r}") from exc
    if km < 0:
        raise ValueError(f"Starting Odometer can't be negative: {val!r}")
    return km


def read_and_clean_vehicle_master(path: Path) -> pd.DataFrame:
    df = pd.read_excel(path, sheet_name=SHEET_NAME)
    df.columns = [str(c).strip() for c in df.columns]

    missing = set(_COLUMN_RENAMES) - set(df.columns)
    if missing:
        raise ValueError(f"'{path.name}' is missing expected columns: {missing}")

    for col in _OPTIONAL_COLUMN_RENAMES:
        if col not in df.columns:
            df[col] = None
    df = df.rename(columns={**_COLUMN_RENAMES, **_OPTIONAL_COLUMN_RENAMES})
    df["base_license_plate"] = df["base_license_plate"].apply(_clean_str)
    df["customer_name"] = df["customer_name"].apply(_clean_str)
    df["oem"] = df["oem"].apply(clean_oem)
    df["vehicle_type"] = df["vehicle_type"].apply(_clean_str)
    df["vehicle_model"] = df["vehicle_model"].apply(_clean_str)
    df["device_installation_date"] = df["device_installation_date"].apply(_parse_install_date)
    df["starting_odometer"] = df["starting_odometer"].apply(_parse_odometer)

    if df["base_license_plate"].isna().any():
        raise ValueError(f"'{path.name}' has a row with a blank Vehicle Number")

    return df[FINAL_COLUMN_ORDER].reset_index(drop=True)


def write_cleaned_copy(path: Path, cleaned_df: pd.DataFrame) -> None:
    """Rewrites the sheet's data rows in place with the cleaned values --
    same headers/column order/Notes sheet, just standardized OEM text and
    install dates as real, consistently-formatted Excel dates."""
    from openpyxl import load_workbook

    all_renames = {**_COLUMN_RENAMES, **_OPTIONAL_COLUMN_RENAMES}
    display_df = cleaned_df.rename(columns={v: k for k, v in all_renames.items()})
    display_df = display_df[list(all_renames.keys())]
    display_df = display_df.sort_values(["Customer Name", "Vehicle Number"]).reset_index(drop=True)

    wb = load_workbook(path)
    ws = wb[SHEET_NAME]

    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    for c, header in enumerate(all_renames.keys(), start=1):
        ws.cell(row=1, column=c, value=header)

    date_col = list(_COLUMN_RENAMES.keys()).index("Device Installation Date") + 1
    for r, row in enumerate(display_df.itertuples(index=False), start=2):
        for c, val in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=val)
            if c == date_col and val is not None:
                cell.number_format = INSTALL_DATE_NUMBER_FORMAT

    wb.save(path)
